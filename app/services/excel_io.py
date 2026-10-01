from io import BytesIO
from typing import BinaryIO

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from app.core.config import TEMPLATE_HEADERS
from app.schemas.reconciliation import ComparisonRow, DuplicateStrategy, ProductRow


class ExcelParseError(Exception):
    def __init__(self, message: str, row: int | None = None):
        self.row = row
        super().__init__(message)


def _normalize_header(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip().lower()


def _expected_header_map() -> dict[str, str]:
    return {_normalize_header(h): h for h in TEMPLATE_HEADERS}


def parse_product_excel(
    source: BinaryIO,
    source_label: str = "file",
    duplicate_strategy: DuplicateStrategy = DuplicateStrategy.REMOVE,
) -> list[ProductRow]:
    raw_rows = _parse_all_product_rows(source, source_label)
    return apply_duplicate_strategy(raw_rows, duplicate_strategy)


def _parse_all_product_rows(source: BinaryIO, source_label: str = "file") -> list[ProductRow]:
    try:
        wb = load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:
        raise ExcelParseError(f"Cannot read {source_label}: not a valid .xlsx file") from exc

    ws = wb.active
    if ws is None:
        raise ExcelParseError(f"{source_label}: workbook has no active sheet")

    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        raise ExcelParseError(f"{source_label}: sheet is empty")

    header_map = _expected_header_map()
    col_index: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        key = _normalize_header(cell)
        if key in header_map and header_map[key] not in col_index:
            col_index[header_map[key]] = idx

    missing = [h for h in TEMPLATE_HEADERS if h not in col_index]
    if missing:
        raise ExcelParseError(
            f"{source_label}: missing columns {', '.join(missing)}. "
            f"Required: {', '.join(TEMPLATE_HEADERS)}"
        )

    products: list[ProductRow] = []
    excel_row_num = 2

    for row in rows_iter:
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            excel_row_num += 1
            continue

        def cell(col_name: str):
            i = col_index[col_name]
            return row[i] if i < len(row) else None

        sku_raw = cell("SKU")
        if sku_raw is None or str(sku_raw).strip() == "":
            raise ExcelParseError(f"{source_label}: row {excel_row_num}: SKU is required", excel_row_num)

        sku = str(sku_raw).strip()

        name_raw = cell("Nama Product")
        name = str(name_raw).strip() if name_raw is not None else ""

        stock_raw = cell("Stok")
        if stock_raw is None or str(stock_raw).strip() == "":
            raise ExcelParseError(f"{source_label}: row {excel_row_num}: Stok is required", excel_row_num)
        try:
            stock = int(stock_raw) if not isinstance(stock_raw, bool) else int(stock_raw)
        except (TypeError, ValueError) as exc:
            raise ExcelParseError(
                f"{source_label}: row {excel_row_num}: Stok must be an integer",
                excel_row_num,
            ) from exc

        no_val = cell("No")
        no: int | None = None
        if no_val is not None and str(no_val).strip() != "":
            try:
                no = int(no_val)
            except (TypeError, ValueError):
                no = None

        products.append(ProductRow(no=no, sku=sku, name=name, stock=stock))
        excel_row_num += 1

    wb.close()
    if not products:
        raise ExcelParseError(f"{source_label}: no data rows found")

    return products


def apply_duplicate_strategy(
    rows: list[ProductRow],
    strategy: DuplicateStrategy,
) -> list[ProductRow]:
    if strategy == DuplicateStrategy.MERGE:
        merged: dict[str, ProductRow] = {}
        for row in rows:
            if row.sku not in merged:
                merged[row.sku] = row.model_copy()
            else:
                current = merged[row.sku]
                merged[row.sku] = ProductRow(
                    no=current.no,
                    sku=row.sku,
                    name=current.name or row.name,
                    stock=current.stock + row.stock,
                )
        return list(merged.values())

    seen: set[str] = set()
    unique: list[ProductRow] = []
    for row in rows:
        if row.sku in seen:
            continue
        seen.add(row.sku)
        unique.append(row)
    return unique


def duplicate_stats_from_check(result) -> tuple[int, int]:
    sku_count = len(result.app) + len(result.real)
    extra_rows = 0
    for entry in result.app + result.real:
        extra_rows += max(0, len(entry.row_numbers) - 1)
    return sku_count, extra_rows


def find_duplicate_skus(source: BinaryIO, source_label: str = "file") -> list[tuple[str, list[int]]]:
    """Return SKUs that appear on more than one Excel row (1-based row numbers)."""
    try:
        wb = load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:
        raise ExcelParseError(f"Cannot read {source_label}: not a valid .xlsx file") from exc

    ws = wb.active
    if ws is None:
        raise ExcelParseError(f"{source_label}: workbook has no active sheet")

    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        raise ExcelParseError(f"{source_label}: sheet is empty")

    header_map = _expected_header_map()
    col_index: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        key = _normalize_header(cell)
        if key in header_map and header_map[key] not in col_index:
            col_index[header_map[key]] = idx

    missing = [h for h in TEMPLATE_HEADERS if h not in col_index]
    if missing:
        raise ExcelParseError(
            f"{source_label}: missing columns {', '.join(missing)}. "
            f"Required: {', '.join(TEMPLATE_HEADERS)}"
        )

    sku_col = col_index["SKU"]
    sku_rows: dict[str, list[int]] = {}
    excel_row_num = 2

    for row in rows_iter:
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            excel_row_num += 1
            continue
        sku_raw = row[sku_col] if sku_col < len(row) else None
        if sku_raw is None or str(sku_raw).strip() == "":
            excel_row_num += 1
            continue
        sku = str(sku_raw).strip()
        sku_rows.setdefault(sku, []).append(excel_row_num)
        excel_row_num += 1

    wb.close()
    return [(sku, rows) for sku, rows in sorted(sku_rows.items()) if len(rows) > 1]


def build_template_workbook() -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Product Stock"
    for col, header in enumerate(TEMPLATE_HEADERS, start=1):
        ws.cell(row=1, column=col, value=header)
        ws.column_dimensions[get_column_letter(col)].width = 18
    ws.append([1, "SKU-001", "Contoh Product", 100])
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def build_corrected_workbook(rows: list[ComparisonRow]) -> BytesIO:
    from app.services.reconcile import build_export_rows

    export_rows = build_export_rows(rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "Corrected Stock"
    for col, header in enumerate(TEMPLATE_HEADERS, start=1):
        ws.cell(row=1, column=col, value=header)

    for idx, row in enumerate(export_rows, start=1):
        ws.append([idx, row.sku, row.name, row.stock])

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
